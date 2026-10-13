"""Build common editable figure objects from one source-derived pixel plan.

No source inference or automatic visual approval. Unusual curves/MathType remain
scoped construction work; the existing figure_workflow owns review/publication.
"""
import argparse
from io import BytesIO
import json
import math
import os
from pathlib import Path
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from figure_workflow import Workflow, task_lock, read, require, APP
from backend.layout import project_root
from backend.safety import atomic_json, digest, DomainError

PAGES = {'wide': (960, 540), 'standard': (720, 540),
         'a4_landscape': (841.89, 595.28)}
PRESETS = {'rect': MSO_SHAPE.RECTANGLE, 'rounded_rect': MSO_SHAPE.ROUNDED_RECTANGLE,
           'ellipse': MSO_SHAPE.OVAL, 'diamond': MSO_SHAPE.DIAMOND,
           'triangle': MSO_SHAPE.ISOSCELES_TRIANGLE, 'right_arrow': MSO_SHAPE.RIGHT_ARROW,
           'down_arrow': MSO_SHAPE.DOWN_ARROW, 'parallelogram': MSO_SHAPE.PARALLELOGRAM}
LAYERS = ['background', 'path', 'node', 'complex_image', 'icon', 'text']
ARROWS = {'none', 'triangle', 'stealth', 'diamond', 'oval', 'arrow'}


def numbers(value, count):
    require(isinstance(value, (list, tuple)) and len(value) == count and
            all(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)
                for v in value), 'figure_plan_geometry_invalid')
    return value


def color(value):
    require(isinstance(value, str) and len(value.lstrip('#')) == 6, 'figure_plan_color_invalid')
    return RGBColor.from_string(value.lstrip('#'))


def style(shape, item, scale):
    fill = item.get('fill')
    if hasattr(shape, 'fill'):
        if fill is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = color(fill)
    stroke = item.get('stroke')
    if stroke is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(float(item.get('line_width_pt', 1)) * scale)
        if item.get('dash'):
            shape.line.dash_style = {'dash': MSO_LINE_DASH_STYLE.DASH,
                                     'dot': MSO_LINE_DASH_STYLE.ROUND_DOT}[item['dash']]
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    for field, tag in [('begin_arrow', 'a:headEnd'), ('end_arrow', 'a:tailEnd')]:
        arrow = item.get(field, 'none')
        require(arrow in ARROWS, 'figure_plan_arrow_invalid')
        if arrow != 'none':
            element = OxmlElement(tag)
            element.set('type', arrow)
            element.set('w', 'sm')
            element.set('len', 'sm')
            shape._element.spPr.get_or_add_ln().append(element)


def text(shape, item, scale):
    frame = shape.text_frame
    frame.clear()
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = bool(item.get('word_wrap', False))
    size = float(item['font_size_pt'])
    require(1 <= size <= 200, 'figure_plan_font_size_required')
    margin = float(item.get('margin_pt', 0 if item['kind'] == 'text' else size * .2)) * scale
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Pt(margin)
    frame.vertical_anchor = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE,
                             'bottom': MSO_ANCHOR.BOTTOM}[item.get('valign', 'middle')]
    for i, line in enumerate(item.get('text', '').split('\n')):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
                       'right': PP_ALIGN.RIGHT}[item.get('align', 'center')]
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(float(item.get('line_spacing_pt', size)) * scale)
        p.font.name = item.get('font', 'Times New Roman')
        p.font.size = Pt(size * scale)
        p.font.bold = bool(item.get('bold', False))
        p.font.italic = bool(item.get('italic', False))
        p.font.color.rgb = color(item.get('color', '000000'))
        if item.get('small_caps'):
            p._p.get_or_add_pPr().get_or_add_defRPr().set('cap', 'small')
            for run in p.runs:
                run._r.get_or_add_rPr().set('cap', 'small')


def bezier_geometry(shape, commands, point, closed=False):
    """One native freeform, never custom geometry inside a connector container."""
    require(isinstance(commands, list) and commands and isinstance(commands[0],list)
            and commands[0] and commands[0][0] == 'M',
            'figure_plan_geometry_invalid')
    segments, points = [], []
    for command in commands:
        require(isinstance(command, list) and command and command[0] in {'M', 'L', 'C', 'Q'},
                'figure_plan_geometry_invalid')
        count = {'M': 2, 'L': 2, 'C': 6, 'Q': 4}[command[0]]
        values = numbers(command[1:], count)
        converted = [point(values[n:n+2]) for n in range(0, count, 2)]
        points.extend(converted)
        segments.append((command[0], converted))
    left, top = min(p[0] for p in points), min(p[1] for p in points)
    width = max(1, max(p[0] for p in points) - left)
    height = max(1, max(p[1] for p in points) - top)
    shape.left, shape.top, shape.width, shape.height = left, top, width, height
    props = shape._element.spPr
    props.xfrm.attrib.clear()
    for old in list(props):
        if old.tag.endswith('}prstGeom') or old.tag.endswith('}custGeom'):
            props.remove(old)
    geometry = OxmlElement('a:custGeom')
    for tag in ('avLst', 'gdLst', 'ahLst', 'cxnLst'):
        geometry.append(OxmlElement('a:' + tag))
    rect = OxmlElement('a:rect')
    for key, value in dict(l='l', t='t', r='r', b='b').items(): rect.set(key, value)
    geometry.append(rect)
    paths, path = OxmlElement('a:pathLst'), OxmlElement('a:path')
    path.set('w', str(width)); path.set('h', str(height)); path.set('fill', 'norm' if closed else 'none')
    for kind, coordinates in segments:
        segment = OxmlElement('a:' + {'M': 'moveTo', 'L': 'lnTo', 'C': 'cubicBezTo', 'Q': 'quadBezTo'}[kind])
        for x, y in coordinates:
            node = OxmlElement('a:pt'); node.set('x', str(x-left)); node.set('y', str(y-top)); segment.append(node)
        path.append(segment)
    if closed: path.append(OxmlElement('a:close'))
    paths.append(path); geometry.append(paths)
    props.insert(1, geometry)


def connector_route(shape, item):
    points = item.get('route_points', [])
    if not points: return
    require(3 <= len(points) <= 6, 'figure_plan_geometry_invalid')
    for value in points: numbers(value, 2)
    sx, sy = points[0]; ex, ey = points[-1]
    geometry = shape._element.spPr.prstGeom
    geometry.set('prst', 'bentConnector'+str(len(points)-1))
    adjustments = geometry.avLst
    for old in list(adjustments): adjustments.remove(old)
    if len(points) >= 4:
        require(ex != sx, 'figure_plan_geometry_invalid')
        node = OxmlElement('a:gd'); node.set('name','adj1'); node.set('fmla','val '+str(round((points[1][0]-sx)/(ex-sx)*100000))); adjustments.append(node)
    if len(points) >= 5:
        require(ey != sy, 'figure_plan_geometry_invalid')
        node = OxmlElement('a:gd'); node.set('name','adj2'); node.set('fmla','val '+str(round((points[2][1]-sy)/(ey-sy)*100000))); adjustments.append(node)
    line = shape._element.spPr.get_or_add_ln()
    if len(points) == 6:
        node = OxmlElement('a:gd'); node.set('name','adj3'); node.set('fmla','val '+str(round((points[3][0]-sx)/(ex-sx)*100000))); adjustments.append(node)
    line.insert_element_before(OxmlElement('a:round'), 'a:headEnd', 'a:tailEnd', 'a:extLst')


def build(workflow, plan_path):
    plan_path = workflow.owned(plan_path)
    plan = read(plan_path)
    require(plan.get('source_image_sha256') == digest(workflow.source), 'artifact_hash_mismatch')
    require(not (workflow.output / 'result.json').exists(), 'figure_result_already_published')
    require(len(workflow.state['rounds']) < 2, 'figure_review_limit_reached')
    items = plan.get('objects')
    require(isinstance(items, list) and 2 <= len(items) <= 1000, 'figure_plan_objects_required')
    selection = workflow.context['figure_selection']
    page_w, page_h = PAGES[selection['page_size']]
    src_w, src_h = numbers(selection['source_size_pt'], 2)
    require(src_w > 0 and src_h > 0, 'figure_plan_geometry_invalid')
    with Image.open(workflow.source) as source:
        source.load()
        source = source.copy()
    image_w, image_h = source.size
    require([image_w, image_h] == [selection['image']['width'], selection['image']['height']],
            'figure_source_size_mismatch')
    scale = min(page_w / src_w, page_h / src_h) * .95 if selection.get('allow_uniform_scale') else 1
    require(src_w * scale <= page_w and src_h * scale <= page_h, 'figure_source_does_not_fit_page')
    kx, ky = src_w / image_w * scale, src_h / image_h * scale
    ox, oy = (page_w - src_w * scale) / 2, (page_h - src_h * scale) / 2

    def point(value):
        x, y = numbers(value, 2)
        require(0 <= x <= image_w and 0 <= y <= image_h, 'figure_plan_outside_source')
        return Pt(ox + x * kx), Pt(oy + y * ky)

    def box(value):
        x, y, w, h = numbers(value, 4)
        require(w > 0 and h > 0, 'figure_plan_geometry_invalid')
        point((x + w, y + h))
        return *point((x, y)), Pt(w * kx), Pt(h * ky)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(page_w), Pt(page_h)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(plan.get('background', 'FFFFFF'))
    registry, expected_text, equations, text_objects, shapes, ordered = [], [], [], [], {}, []
    defaults = plan.get('defaults', {})
    styles = plan.get('styles', {})
    for original in items:
        item = {**defaults, **styles.get(original.get('style'), {}), **original}
        # Old plans used a straight connector as a placeholder for a source
        # Bezier. PowerPoint rejects some custom geometries inside p:cxnSp.
        # These precise geometric arcs are native freeforms; ordinary module
        # connectors continue to use actual bound Office connectors.
        if item.get('native_bezier'):
            item['kind'] = 'bezier'
        name, kind = item['name'], item['kind']
        require(isinstance(name, str) and 0 < len(name) <= 120 and name not in shapes,
                'figure_plan_object_name_invalid')
        layer = item.get('layer', 'path' if kind in {'connector', 'line', 'polyline', 'bezier'} else
                         'complex_image' if kind == 'image' else 'text' if kind in {'text', 'formula'} else 'node')
        require(layer in LAYERS, 'figure_plan_layer_invalid')
        if kind in PRESETS:
            shape = slide.shapes.add_shape(PRESETS[kind], *box(item['bbox']))
            values = item.get('adjustments', [item['adjustment']] if 'adjustment' in item else [])
            require(isinstance(values, list) and len(values) <= len(shape.adjustments), 'figure_plan_adjustments_invalid')
            numbers(values, len(values))
            for index, value in enumerate(values): shape.adjustments[index] = value
        elif kind in {'text', 'formula'}:
            shape = slide.shapes.add_textbox(*box(item['bbox']))
        elif kind == 'connector':
            connector_kind = {'straight': MSO_CONNECTOR.STRAIGHT, 'elbow': MSO_CONNECTOR.ELBOW, 'curved': MSO_CONNECTOR.CURVE}[item.get('routing', 'straight')]
            shape = slide.shapes.add_connector(connector_kind, *point(item['start']), *point(item['end']))
        elif kind == 'bezier':
            shape = slide.shapes.add_textbox(Pt(0), Pt(0), Pt(1), Pt(1))
            shape._element.nvSpPr.cNvSpPr.attrib.pop('txBox', None)
            bezier_geometry(shape, item.get('native_bezier', item.get('commands')), point, bool(item.get('closed_fill', False)))
        elif kind in {'line', 'polyline', 'polygon'}:
            points = item.get('points') or [item['start'], item['end']]
            require(len(points) >= (3 if kind == 'polygon' else 2), 'figure_plan_geometry_invalid')
            points = [point(p) for p in points]
            freeform = slide.shapes.build_freeform(*points[0])
            freeform.add_line_segments(points[1:], close=kind == 'polygon')
            shape = freeform.convert_to_shape()
        elif kind == 'image':
            # Independent image assets may already be cropped/masked. Otherwise
            # crop a named complex-image region; never use the whole source as a backdrop.
            if item.get('asset'):
                asset = workflow.owned(item['asset'])
                require(digest(asset) == item.get('asset_sha256'), 'artifact_hash_mismatch')
                with Image.open(asset) as im:
                    pic = im.copy()
            else:
                crop = numbers(item.get('crop', item['bbox']), 4)
                x, y, w, h = crop
                box(crop)
                require(w * h < image_w * image_h * .95, 'figure_full_source_image_forbidden')
                pic = source.crop((round(x), round(y), round(x + w), round(y + h)))
            _, _, w, h = numbers(item['bbox'], 4)
            require(abs((w * kx / (h * ky)) / (pic.width / pic.height) - 1) <= .02,
                    'figure_image_aspect_mismatch')
            data = BytesIO()
            pic.save(data, 'PNG')
            data.seek(0)
            shape = slide.shapes.add_picture(data, *box(item['bbox']))
        else:
            raise DomainError('figure_plan_unsupported_object', kind=kind)
        shape.name = name
        if 'rotation' in item:
            require(kind != 'connector', 'figure_plan_rotate_connector_invalid')
            shape.rotation = numbers([item['rotation']], 1)[0] % 360
        shapes[name] = shape
        if kind != 'image':
            style(shape, item, scale)
        if 'text' in item:
            require(kind in PRESETS or kind == 'text', 'figure_plan_text_kind_invalid')
            text(shape, item, scale)
            expected_text.append(dict(name=name, text=item['text']))
            text_objects.append(dict(name=name, resize_frame=kind == 'text'))
        if kind == 'formula':
            require(selection['formula_mode'] == 'office', 'figure_plan_requires_special_formula_pipeline')
            linear = item.get('native_linear')
            require(isinstance(linear, str) and linear.strip(), 'figure_plan_formula_required')
            text(shape, item, scale)
            equations.append(dict(id=name, native_linear=linear, font_size=float(item['font_size_pt']) * scale,
              target=dict(type='powerpoint_shape', slide_index=1, shape_name=name)))
        registry.append(dict(name=name, layer=layer))
        ordered.append((layer, item, shape))

    # Bind after all nodes exist, then restore source endpoints; node centres are
    # not an excuse to replace faithful off-centre source connections.
    for layer, item, shape in ordered:
        if item['kind'] != 'connector':
            continue
        for field, bind in [('source', shape.begin_connect), ('target', shape.end_connect)]:
            endpoint = item[field]
            require(endpoint['name'] in shapes and endpoint['name'] != item['name'],
                    'figure_plan_connector_target_missing')
            bind(shapes[endpoint['name']], int(endpoint['site']))
        shape.begin_x, shape.begin_y = point(item['start'])
        shape.end_x, shape.end_y = point(item['end'])
        connector_route(shape, item)
    for layer, item, shape in sorted(ordered, key=lambda row: LAYERS.index(row[0])):
        slide.shapes._spTree.remove(shape._element)
        slide.shapes._spTree.insert_element_before(shape._element, 'p:extLst')

    axes = plan.get('axis_constraints', [])
    for axis in axes:
        require(axis.get('source') in shapes and axis.get('target') in shapes and
                axis.get('axis') in {'horizontal', 'vertical'}, 'figure_plan_axis_invalid')
    manifest = dict(slide_index=1, expected_slide_count=1, render_width=1920,
      render_height=round(1920 * page_h / page_w), delivery_gate=True, require_full_registry=True,
      expected_shape_count=len(items), expected_image_count=sum(i['kind'] == 'image' for i in items),
      expected_connector_count=sum(item['kind'] == 'connector' for _, item, _ in ordered),
      expected_formula_count=len(equations), expected_axis_constraint_count=len(axes),
      objects=registry, expected_text=expected_text,
      formula_objects=[dict(name=e['id'], kind='office') for e in equations], axis_constraints=axes)
    # Atomic working-file replacement only after every object/asset is valid.
    # Never touch target, registered artifacts, existing review rounds or source.
    staging = workflow.owned(workflow.working.with_suffix('.building.pptx'), exists=False)
    try:
        prs.save(staging)
        os.replace(staging, workflow.working)
    finally:
        if staging.exists():
            staging.unlink()
    audit_path = workflow.output / 'native-audit-input.json'
    atomic_json(audit_path, manifest)
    equations_path = workflow.output / 'native-equations-input.json'
    if equations or text_objects:
        atomic_json(equations_path, dict(mode='Office', equations=equations, text_objects=text_objects))
    return dict(working_path=workflow.relative(workflow.working), manifest_path=workflow.relative(audit_path),
                equation_manifest_path=workflow.relative(equations_path) if equations else None,
                preparation_manifest_path=workflow.relative(equations_path) if equations or text_objects else None,
                objects=len(items), source_plan_sha256=digest(plan_path))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('context')
    parser.add_argument('plan')
    args = parser.parse_args()
    workflow = Workflow(project_root(APP), args.context)
    with task_lock(workflow.output / 'figure-workflow.lock'):
        workflow = Workflow(project_root(APP), args.context)
        workflow.live()
        result = build(workflow, args.plan)
        if result['preparation_manifest_path']:
            workflow.request('/jobs/' + workflow.job_id + '/office-native',
              dict(operation='prepare', presentation=result['working_path'],
                   manifest=result['preparation_manifest_path']))
        workflow.save('批量构建完成；运行一次合并审查')
        print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    main()
